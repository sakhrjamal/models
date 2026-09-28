# -*- coding: utf-8 -*-
"""زرع البيانات المرجعية في قاعدة البيانات"""
import sqlite3, re, os, sys, glob

HERE = os.path.dirname(os.path.abspath(__file__))
DB   = os.environ.get('QMS_DB') or os.path.join(HERE, 'data', 'qms.db')


def _find_index_file():
    exact = os.path.join(HERE, 'data', 'فهرس_الاكواد_11_11_2023.xlsm')
    if os.path.exists(exact):
        return exact
    # احتياط: أي ملف اكسل في مجلد data (قد يختلف اسمه بعد فك الضغط)
    hits = glob.glob(os.path.join(HERE, 'data', '*.xlsm')) + glob.glob(os.path.join(HERE, 'data', '*.xlsx'))
    return hits[0] if hits else exact


SRC = _find_index_file()

MACHINES = [
 ('SL-01','الأسليتر','القص الطولي','SL',None),
 ('FD-05','ماكينة الطي 5','التقطيع والطي','F','5x5 cm'),
 ('FD-75','ماكينة الطي 7.5','التقطيع والطي','S','7.5x7.5 cm'),
 ('FD-10','ماكينة الطي 10','التقطيع والطي','T','10x10 cm'),
 ('PK-01','ماكينة التغليف','التغليف',None,None),
 ('EO-01','ماكينة التعقيم','التعقيم','EO',None),
 ('LN-SC','خط الشاش المقطع','تعبئة أو تغليف','C',None),
 ('LN-LS','خط اللاب سبونج','التغليف','L',None),
 ('LN-SR','خط الرول الصغير','التغليف','R',None),
]
MATRIX = [
 ('FD-05-8','FD-05',8,13,'5x5 cm'),   ('FD-05-12','FD-05',12,18,'5x5 cm'),
 ('FD-05-16','FD-05',16,23,'5x5 cm'), ('FD-75-8','FD-75',8,18,'7.5x7.5 cm'),
 ('FD-75-12','FD-75',12,26.5,'7.5x7.5 cm'),('FD-75-16','FD-75',16,33,'7.5x7.5 cm'),
 ('FD-10-8','FD-10',8,23,'10x10 cm'), ('FD-10-12','FD-10',12,33,'10x10 cm'),
 ('FD-10-16','FD-10',16,44,'10x10 cm'),
]
MACHINE_CUT = [
 ('FD-05',20,'طول القطعة المفرودة على ماكينة 5 — ثابت أيًا كان العرض'),
 ('FD-75',30,'طول القطعة المفرودة على ماكينة 7.5'),
 ('FD-10',40,'طول القطعة المفرودة على ماكينة 10'),
]
MONTHS = [('JAN','يناير',1),('FEB','فبراير',2),('MAR','مارس',3),('APR','أبريل',4),
          ('MAY','مايو',5),('JUN','يونيو',6),('JUL','يوليو',7),('AUG','أغسطس',8),
          ('SEP','سبتمبر',9),('OCT','أكتوبر',10),('NOV','نوفمبر',11),('DEC','ديسمبر',12)]
SETTINGS = [
 ('aeration_min_h','24','المدة الدنيا للتهوية بالساعة — تُضبط من تقرير التأهيل'),
 ('aeration_temp_min','18','أدنى حرارة لغرفة التهوية'),
 ('aeration_temp_max','35','أعلى حرارة لغرفة التهوية'),
 ('scrap_limit','0.05','حد قبول الفاقد في الطي'),
 ('scrap_limit_machine','0.03','حد قبول التالف في سجل المكائن'),
 ('slitter_edge_trim_each_cm','1','تنظيف ثابت من كل طرف للرول قبل احتساب العرض التشغيلي للأسليتر بالسم'),
 ('slitter_warning_limit_cm','4','حد مرجعي لفرق خطة القص؛ تجاوزه يظهر تحذيرًا إضافيًا ولا يمنع الحفظ'),
 ('slitter_max_trim_cm','4','إعداد قديم محفوظ للتوافق — لم يعد يمنع الحفظ'),
 ('wip_max_days','30','المدة القصوى لانتظار الكرتونة الوسيطة'),
 ('shelf_life_m','60','مدة الصلاحية الافتراضية بالأشهر'),
 ('factory_name','مصنع الشاش والأربطة الطبية — الأحساء',''),
 ('default_sr_length_m','2000','الطول النظري للرول للتخطيط — الفعلي يُدخل في سند الأسليتر'),
 ('default_jumbo_width_cm','90','عرض الجامبو الافتراضي للتخطيط'),
 ('yield_loss_pct','0','معامل هدر الإنتاج % — يُضبط لاحقًا من واقع الأوزان'),
 ('gsm','17','وزن الشاش جم/م² للحساب النظري'),
 ('slitter_edge_trim_each_cm','1','تنظيف كل طرف من الجامبو (سم)'),
 ('slitter_warning_limit_cm','4','حد التحذير لفرق خطة القص (سم)'),
 ('backup_interval_h','6','الفاصل بين النسخ الاحتياطية التلقائية بالساعة'),
 ('backup_keep','60','عدد النسخ الاحتياطية المحفوظة (تُحذف الأقدم)'),
 ('session_hours','12','مدة بقاء جلسة تسجيل الدخول بالساعات'),
]
TY = {'RR':('مادة خام','Raw Material'),'PBT':('مادة خام','Raw Material'),
      'SR':('نصف مصنع','WIP'),'SP':('نصف مصنع','WIP'),
      'GS':('منتج تام','Finished Good'),'GB':('منتج تام','Finished Good'),
      'PK':('مواد تعبئة — غير معقم','Packaging - Non Sterile'),
      'BX':('مواد تعبئة — معقم','Packaging - Sterile'),
      'MB':('مواد تعبئة — كرتون خارجي','Master Box')}
UOM = {'RR':'رول','PBT':'رول','SR':'رول','SP':'قطعة','GS':'قطعة','GB':'قطعة',
       'PK':'قطعة','BX':'قطعة','MB':'قطعة'}
MC = {'5x5 cm':'FD-05','7.5x7.5 cm':'FD-75','10x10 cm':'FD-10'}

def pf(c):
    m = re.match(r'^([A-Za-z]+)', c)
    return m.group(1).upper() if m else '?'

def parse(d, code):
    o = {}
    m = re.search(r'(\d+(?:\.\d+)?)\s*[×xX*]\s*(\d+(?:\.\d+)?)\s*CM', d, re.I)
    o['size'] = f"{m.group(1)}x{m.group(2)} cm" if m else None
    m = re.search(r'(\d+)\s*PLY', d, re.I); o['ply'] = int(m.group(1)) if m else None
    o['ster'] = 'غير معقم' if re.search(r'NON-?\s*STERILE', d, re.I) else (
                'معقم' if re.search(r'STERILE', d, re.I) else None)
    o['xray'] = 'WITH X-RAY' if re.search(r'X-?\s*RAY', d, re.I) else (
                'WITHOUT X-RAY' if re.search(r'Plain', d, re.I) else None)
    if re.search(r'19\s*[×xX*]\s*15', d):   o['mesh'] = '19x15'
    elif re.search(r'18\s*[×xX*]\s*11', d): o['mesh'] = '18x11'
    else: o['mesh'] = '19x15' if code.endswith('M') else '18x11'
    o['edge'] = 'Folded Edge' if re.search(r'FOLD', d, re.I) else None
    return o

def load_items():
    import openpyxl
    wb = openpyxl.load_workbook(SRC, data_only=True)
    s  = wb['فهرس الاكواد النهائي 7-10-2023']
    raw = []
    for r in range(3, 304):
        c = s.cell(r, 2).value
        if not c: continue
        raw.append(dict(code=str(c).strip().upper(), fam=s.cell(r,3).value,
                        contain=s.cell(r,5).value, wide=s.cell(r,6).value,
                        length=s.cell(r,7).value, desc=str(s.cell(r,8).value or '')))
    pack, mbx = {}, {}
    for r in raw:
        p = pf(r['code'])
        if p not in ('PK','BX','MB'): continue
        for g in re.findall(r'\b(GS\d+M?|GB\d+)\b', r['desc']):
            (mbx if p == 'MB' else pack).setdefault(g, []).append(r['code'])
    rows = []
    for r in raw:
        p = pf(r['code']); a = parse(r['desc'], r['code'])
        if p not in ('GS','SP','RR','SR','GB','PBT'): a['mesh'] = None
        key = None
        if p in ('GS','SP') and all([a['size'], a['ply'], a['xray'], a['mesh'], a['ster']]):
            key = f"{p}|{a['size']}|{a['ply']}|{a['xray']}|{a['mesh']}|{a['ster']}"
        rows.append((
            r['code'], p, TY.get(p, ('',''))[0], TY.get(p, ('',''))[1], r['fam'],
            a['size'], a['ply'], a['xray'], a['mesh'], a['ster'], a['edge'],
            MC.get(a['size']) if p in ('GS','SP') else None, UOM.get(p, ''),
            r['wide'] if p in ('RR','SR','PBT') else None,
            r['length'] if p in ('RR','SR','PBT') else None,
            r['contain'] if isinstance(r['contain'], str) else None,
            ', '.join(dict.fromkeys(pack.get(r['code'], []))) or None,
            ', '.join(dict.fromkeys(mbx.get(r['code'], []))) or None,
            r['desc'], key, 'نشط'))
    return rows

def main():
    os.makedirs(os.path.dirname(DB), exist_ok=True)
    con = sqlite3.connect(DB)
    con.executescript(open(os.path.join(HERE, 'schema.sql'), encoding='utf-8').read())
    con.executemany('INSERT OR REPLACE INTO machines VALUES(?,?,?,?,?,1)', MACHINES)
    con.executemany('INSERT OR REPLACE INTO slit_matrix VALUES(?,?,?,?,?)', MATRIX)
    con.executemany('INSERT OR REPLACE INTO months VALUES(?,?,?)', MONTHS)
    con.executemany('INSERT OR REPLACE INTO machine_cut VALUES(?,?,?)', MACHINE_CUT)
    con.executemany('INSERT OR REPLACE INTO settings VALUES(?,?,?)', SETTINGS)
    if os.path.exists(SRC):
        rows = load_items()
        con.executemany('INSERT OR REPLACE INTO items VALUES(' + ','.join(['?']*21) + ')', rows)
        print(f'  الأصناف     : {len(rows)}')
    con.commit()
    for t in ('items','machines','slit_matrix','months','settings'):
        n = con.execute(f'SELECT COUNT(*) FROM {t}').fetchone()[0]
        print(f'  {t:14}: {n}')
    con.close()
    print('\nقاعدة البيانات:', DB)

if __name__ == '__main__':
    main()
